"""Xuất file Excel danh sách đặt hàng theo đúng format mẫu + chuẩn thương hiệu MGV."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from engine import OUTPUT_COLUMNS, ITEM_GROUP_CHOICES

MG_BLUE = "0070C0"
HEADER_FONT = Font(name="Cambria", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11)
HEADER_FILL = PatternFill("solid", fgColor=MG_BLUE)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

# Bề rộng cột (kế thừa từ file mẫu + căn chỉnh hợp lý)
COL_WIDTHS = {
    "No": 5, "Brand": 9, "Item Code": 14, "Item Name": 34, "Item Group": 11,
    "Quantity": 9, "Cost": 12, "PIC": 8, "Analyzer": 24, "Reason of Order": 24,
    "Purpose": 12, "Pending NW": 12, "Onhand HN": 10, "Onhand HCM": 11, "Note": 30,
}
CENTER_COLS = {"No", "Brand", "Item Group", "Quantity", "Cost", "PIC",
               "Pending NW", "Onhand HN", "Onhand HCM"}

# Nhãn khu vực cho 2 cột Min/Max tiêu chuẩn thêm vào (bên phải Note)
REGION_LABEL = {"HN": "HAN", "HCM": "HCM"}


def _region_extra_headers(region):
    label = REGION_LABEL.get(region, region)
    return [f"{label} Min", f"{label} Max"]


def _write_sheet(ws, rows, region):
    extra = _region_extra_headers(region)          # ['HAN Min','HAN Max'] hoặc ['HCM Min','HCM Max']
    all_cols = OUTPUT_COLUMNS + extra
    # Header
    for c, name in enumerate(all_cols, start=1):
        cell = ws.cell(1, c, name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    # Data
    for i, row in enumerate(rows, start=1):
        r = i + 1
        row = dict(row)
        row["No"] = i
        row[extra[0]] = row.get("Std Min")
        row[extra[1]] = row.get("Std Max")
        for c, name in enumerate(all_cols, start=1):
            cell = ws.cell(r, c, row.get(name))
            cell.font = BODY_FONT
            cell.alignment = CENTER if (name in CENTER_COLS or name in extra) else LEFT
            cell.border = BORDER
            if name == "Cost":
                cell.number_format = "#,##0.00"
    # Widths
    for c, name in enumerate(all_cols, start=1):
        w = COL_WIDTHS.get(name, 10 if name in extra else 12)
        ws.column_dimensions[get_column_letter(c)].width = w
    # Freeze header
    ws.freeze_panes = "A2"
    # Dropdown Item Group (cột E)
    if rows:
        col_letter = get_column_letter(OUTPUT_COLUMNS.index("Item Group") + 1)
        dv = DataValidation(
            type="list",
            formula1='"%s"' % ",".join(ITEM_GROUP_CHOICES),
            allow_blank=True, showDropDown=False,
        )
        dv.error = "Chọn một trong: " + ", ".join(ITEM_GROUP_CHOICES)
        dv.prompt = "Mặc định SPP"
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{len(rows) + 1}")


def write_workbook(sheets, path):
    """
    sheets: dict giữ thứ tự { 'HN': [rows...], 'HCM': [rows...] }  (hoặc chỉ 1 khu vực)
    Tên sheet chính là khu vực -> quyết định nhãn 2 cột Min/Max thêm vào.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for region, rows in sheets.items():
        ws = wb.create_sheet(title=region)
        _write_sheet(ws, rows, region)
    wb.save(path)
    return path
